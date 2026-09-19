from __future__ import annotations

import hashlib
import json
import re
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.application.import_contract import CONTRACTS, locate_header, validate_workbook
from app.domain.import_models import (
    CaseLinkStatus,
    CaseRecord,
    JobRecord,
    QCEvidenceStatus,
    TechnicianIdentityStatus,
    TechnicianRecord,
    WorkbookPreview,
)

UNVERIFIED_TECH_CODES = frozenset({"PWS1", "PWS51"})


def _text(value: object) -> str | None:
    if value is None:
        return None
    result = re.sub(r"\s+", " ", str(value)).strip()
    return result or None


def _code(value: object) -> str | None:
    result = _text(value)
    return result.upper() if result else None


def _job_no(value: object) -> str | None:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    result = _text(value)
    return result.upper() if result else None


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        for format_string in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
            try:
                return datetime.strptime(value.strip(), format_string).date()
            except ValueError:
                pass
    return None


def _bool(value: object) -> bool | None:
    if value is None or _text(value) is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    normalised = _text(value).casefold()
    if normalised in {"yes", "y", "true", "1", "มี"}:
        return True
    if normalised in {"no", "n", "false", "0", "ไม่มี"}:
        return False
    return None


def _value(row: tuple, fields: dict[str, int], name: str) -> object:
    index = fields.get(name)
    return row[index] if index is not None and index < len(row) else None


def _serialise(value: object) -> object:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if hasattr(value, "value"):
        return value.value
    return value


def _hash_payload(*values: object) -> str:
    encoded = json.dumps([_serialise(value) for value in values], ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _identity_status(tech_code: str | None, master_ids: set[str]) -> TechnicianIdentityStatus:
    if not tech_code:
        return TechnicianIdentityStatus.MISSING_TECHNICIAN_ID
    if tech_code in UNVERIFIED_TECH_CODES:
        return TechnicianIdentityStatus.UNVERIFIED_TECHNICIAN_IDENTITY
    if tech_code in master_ids:
        return TechnicianIdentityStatus.VERIFIED
    return TechnicianIdentityStatus.PENDING_MASTER_MATCH


def _qc_evidence_status(job_status: str | None, qc_result: str | None) -> QCEvidenceStatus:
    if qc_result:
        return QCEvidenceStatus.RECORDED
    normalised_status = (job_status or "").casefold()
    if normalised_status in {"cancel", "cancelled", "canceled"}:
        return QCEvidenceStatus.NOT_APPLICABLE_CANCELLED
    if normalised_status in {"not complete", "not completed", "incomplete"}:
        return QCEvidenceStatus.NOT_APPLICABLE_NOT_COMPLETE
    return QCEvidenceStatus.MISSING_SOURCE_VALUE


def _sheet_rows(workbook, contract):
    sheet = workbook[contract.sheet_name]
    header_row, _, fields = locate_header(sheet, contract)
    if header_row is None:
        return
    key_field = contract.required_columns[0]
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            if not any(value not in (None, "") for value in row):
                continue
            if _value(row, fields, key_field) in (None, ""):
                continue
            yield row_number, tuple(row), fields


def preview_workbook(path: str | Path, filename: str | None = None) -> WorkbookPreview:
    source = Path(path)
    validations = validate_workbook(source)
    invalid = [item for item in validations if not item.is_valid]
    if invalid:
        reasons = "; ".join(
            f"{item.sheet_name}: missing={','.join(item.missing_columns) or '-'}, blank_keys={item.blank_key_rows}"
            for item in invalid
        )
        raise ValueError(f"Workbook structure is invalid: {reasons}")

    content_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        master_contract, job_contract, case_contract = CONTRACTS[2], CONTRACTS[0], CONTRACTS[1]
        technicians: list[TechnicianRecord] = []
        master_by_id: dict[str, TechnicianRecord] = {}
        duplicate_master_ids: Counter[str] = Counter()
        master_source_rows = 0
        for _, row, fields in _sheet_rows(workbook, master_contract):
            technician_id = _code(_value(row, fields, "technician_id"))
            if not technician_id:
                continue
            master_source_rows += 1
            record = TechnicianRecord(
                technician_id=technician_id,
                technician_name=_text(_value(row, fields, "technician_name")),
                team=_text(_value(row, fields, "team")),
                vendor_name=_text(_value(row, fields, "vendor")),
                source_status=_text(_value(row, fields, "status")),
                area=_text(_value(row, fields, "area")),
            )
            if technician_id in master_by_id:
                duplicate_master_ids[technician_id] += 1
                continue
            master_by_id[technician_id] = record
            technicians.append(record)
        master_ids = set(master_by_id)

        jobs: list[JobRecord] = []
        job_occurrences: Counter[tuple] = Counter()
        tech_codes_by_job: dict[str, set[str]] = defaultdict(set)
        job_status_counts: Counter[str] = Counter()
        identity_counts: Counter[str] = Counter()
        qc_evidence_counts: Counter[str] = Counter()
        qc_result_counts: Counter[str] = Counter()
        job_number_counts: Counter[str] = Counter()
        for row_number, row, fields in _sheet_rows(workbook, job_contract):
            job_no = _job_no(_value(row, fields, "job_no"))
            if not job_no:
                continue
            install_date = _date(_value(row, fields, "install_date"))
            product_model = _text(_value(row, fields, "product_model"))
            source_tech_code = _code(_value(row, fields, "technician_id"))
            job_status = _text(_value(row, fields, "job_status"))
            qc_result = _text(_value(row, fields, "qc_result"))
            complaint = _bool(_value(row, fields, "complaint"))
            rework = _bool(_value(row, fields, "rework"))
            integrity_violation = _bool(_value(row, fields, "integrity_violation"))
            identity_status = _identity_status(source_tech_code, master_ids)
            qc_evidence_status = _qc_evidence_status(job_status, qc_result)
            base_key = (job_no.casefold(), install_date, (product_model or "").casefold())
            job_occurrences[base_key] += 1
            source_key = _hash_payload("JOB", *base_key, job_occurrences[base_key])
            technician_id = source_tech_code if identity_status is TechnicianIdentityStatus.VERIFIED else None
            values = (
                job_no, install_date, product_model, source_tech_code, technician_id, identity_status,
                _text(_value(row, fields, "project")), _text(_value(row, fields, "vendor")),
                _date(_value(row, fields, "qc_date")), qc_result, qc_evidence_status, job_status,
                complaint, rework, integrity_violation,
            )
            jobs.append(JobRecord(source_key, row_number, *values, _hash_payload(*values)))
            if source_tech_code:
                tech_codes_by_job[job_no].add(source_tech_code)
            job_number_counts[job_no] += 1
            job_status_counts[job_status or "(blank)"] += 1
            identity_counts[identity_status.value] += 1
            qc_evidence_counts[qc_evidence_status.value] += 1
            qc_result_counts[qc_result or "(blank)"] += 1

        cases: list[CaseRecord] = []
        case_occurrences: Counter[tuple] = Counter()
        link_counts: Counter[str] = Counter()
        root_cause_counts: Counter[str] = Counter()
        case_status_counts: Counter[str] = Counter()
        complaint_month_counts: Counter[str] = Counter()
        complaint_dates: list[date] = []
        job_dates = [record.install_date for record in jobs if record.install_date]
        job_period_start = min(job_dates) if job_dates else None
        job_period_end = max(job_dates) if job_dates else None
        for row_number, row, fields in _sheet_rows(workbook, case_contract):
            job_no = _job_no(_value(row, fields, "job_no"))
            if not job_no:
                continue
            complaint_date = _date(_value(row, fields, "complaint_date"))
            source_tech_code = _code(_value(row, fields, "technician_id"))
            identity_status = _identity_status(source_tech_code, master_ids)
            technician_id = source_tech_code if identity_status is TechnicianIdentityStatus.VERIFIED else None
            base_key = (job_no.casefold(), complaint_date)
            case_occurrences[base_key] += 1
            source_key = _hash_payload("CASE", *base_key, case_occurrences[base_key])
            completed_date = _date(_value(row, fields, "completed_date"))
            if job_no not in tech_codes_by_job and completed_date and job_period_start and completed_date < job_period_start:
                link_status = CaseLinkStatus.REFERENCE_OUTSIDE_CURRENT_JOB_DATA
            elif job_no not in tech_codes_by_job:
                link_status = CaseLinkStatus.JOB_REFERENCE_REQUIRES_REVIEW
            elif source_tech_code and tech_codes_by_job[job_no] and source_tech_code not in tech_codes_by_job[job_no]:
                link_status = CaseLinkStatus.TECHNICIAN_CONFLICT
            else:
                link_status = CaseLinkStatus.MATCHED
            root_cause_status = _text(_value(row, fields, "root_cause_status"))
            case_status = _text(_value(row, fields, "case_status"))
            values = (
                job_no, completed_date, complaint_date, source_tech_code,
                technician_id, identity_status, _text(_value(row, fields, "issue_category")),
                _text(_value(row, fields, "issue_detail")), _text(_value(row, fields, "qc_result")),
                root_cause_status, _text(_value(row, fields, "root_cause_type")),
                _text(_value(row, fields, "root_cause_detail")), _text(_value(row, fields, "immediate_action")),
                _text(_value(row, fields, "preventive_action")), _bool(_value(row, fields, "rework")),
                _bool(_value(row, fields, "service_mind")), _text(_value(row, fields, "owner")),
                case_status, _date(_value(row, fields, "close_date")), link_status,
            )
            cases.append(CaseRecord(source_key, row_number, *values, _hash_payload(*values)))
            link_counts[link_status.value] += 1
            root_cause_counts[root_cause_status or "(blank)"] += 1
            case_status_counts[case_status or "(blank)"] += 1
            if complaint_date:
                complaint_dates.append(complaint_date)
                complaint_month_counts[complaint_date.strftime("%Y-%m")] += 1

        warnings_list: list[str] = []
        if duplicate_master_ids:
            warnings_list.append(f"Technician Master มี Tech ID ซ้ำ {len(duplicate_master_ids)} รหัส ระบบเก็บรายการแรกหลังปรับช่องว่าง")
        if identity_counts[TechnicianIdentityStatus.UNVERIFIED_TECHNICIAN_IDENTITY.value]:
            warnings_list.append("PWS1/PWS51 ถูกเก็บเป็นปริมาณงาน แต่ไม่สร้างโปรไฟล์หรือคะแนนรายบุคคล")
        if link_counts[CaseLinkStatus.TECHNICIAN_CONFLICT.value]:
            warnings_list.append("เคสที่ Tech ID ขัดกับ Job Data ถูกพักไว้ให้คนตรวจ และไม่นำไปให้คะแนนอัตโนมัติ")
        if link_counts[CaseLinkStatus.REFERENCE_OUTSIDE_CURRENT_JOB_DATA.value]:
            warnings_list.append(
                "งานติดตั้งก่อนช่วง Job Data (เช่น ปี 2025) แต่รับเรื่องร้องเรียนในปี 2026 "
                "ยังนับเป็น Complaint ปี 2026 ตาม Complaint Date และเก็บไว้เป็นหลักฐาน"
            )
        if link_counts[CaseLinkStatus.JOB_REFERENCE_REQUIRES_REVIEW.value]:
            warnings_list.append("เคสที่ควรอยู่ในช่วง Job Data แต่ยังหา Order No. ไม่พบ ถูกพักไว้ให้ตรวจสอบ")

        summary = {
            "technicians": {"source_rows": master_source_rows, "unique_verified": len(technicians), "duplicate_ids": sum(duplicate_master_ids.values())},
            "jobs": {
                "rows": len(jobs), "unique_job_numbers": len(job_number_counts),
                "duplicate_job_numbers": sum(1 for count in job_number_counts.values() if count > 1),
                "period_start": job_period_start.isoformat() if job_period_start else None,
                "period_end": job_period_end.isoformat() if job_period_end else None,
                "status": dict(job_status_counts), "technician_identity": dict(identity_counts),
                "qc_evidence": dict(qc_evidence_counts), "qc_result": dict(qc_result_counts),
            },
            "complaints": {
                "rows": len(cases), "link_status": dict(link_counts),
                "root_cause_status": dict(root_cause_counts), "case_status": dict(case_status_counts),
                "reporting_date_field": "Complaint Date",
                "period_start": min(complaint_dates).isoformat() if complaint_dates else None,
                "period_end": max(complaint_dates).isoformat() if complaint_dates else None,
                "by_month": dict(sorted(complaint_month_counts.items())),
                "missing_complaint_date": len(cases) - len(complaint_dates),
            },
            "evidence_policy": {
                "complaint_truth_source": "Complaint Log", "missing_is_not_pass": True,
                "complaint_reporting_date": "complaint_date",
                "blank_cancel_qc": "NOT_APPLICABLE_CANCELLED",
                "blank_not_complete_qc": "NOT_APPLICABLE_NOT_COMPLETE", "no_delete_on_omission": True,
            },
        }
        return WorkbookPreview(filename or source.name, content_hash, technicians, jobs, cases, summary, warnings_list)
    finally:
        workbook.close()
