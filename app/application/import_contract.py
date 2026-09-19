from __future__ import annotations

import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def normalise_header(value: object) -> str:
    return "".join(str(value or "").strip().casefold().split()).replace(".", "")


@dataclass(frozen=True)
class SheetContract:
    sheet_name: str
    required_columns: tuple[str, ...]
    accepted_aliases: dict[str, tuple[str, ...]]


@dataclass(frozen=True)
class ImportValidation:
    sheet_name: str
    header_row: int | None
    row_count: int
    missing_columns: tuple[str, ...]
    blank_key_rows: int

    @property
    def is_valid(self) -> bool:
        return self.header_row is not None and not self.missing_columns and self.blank_key_rows == 0


CONTRACTS = (
    SheetContract(
        sheet_name="Job Data",
        required_columns=("job_no", "install_date", "technician_id", "job_status"),
        accepted_aliases={
            "job_no": ("Job No.", "Job No", "Job Number", "PO", "Order No"),
            "install_date": ("Install Date", "Completed Date"),
            "technician_id": ("Tech ID", "Technician ID"),
            "product_model": ("Product / Model", "Product", "Model"),
            "project": ("Project",),
            "vendor": ("LSP", "Vendor", "Sub Vendor"),
            "qc_date": ("QC Date",),
            "qc_result": ("QC Result", "QC"),
            "complaint": ("Complaint",),
            "rework": ("Rework",),
            "integrity_violation": ("Integrity violation", "Integrity Violation"),
            "job_status": ("Job Status", "Status"),
        },
    ),
    SheetContract(
        sheet_name="Complaint Log",
        required_columns=("job_no", "complaint_date", "technician_id"),
        accepted_aliases={
            "job_no": ("Job No. (เลขที่ออเดอร์)", "Job No.", "Job No", "PO", "Order No"),
            "completed_date": ("Complete Date (วันที่ติดตั้งสำเร็จ)", "Complete Date"),
            "complaint_date": ("Complaint Date (วันที่ได้รับการร้องเรียน)", "Complaint Date"),
            "technician_id": ("Tech ID", "Technician ID"),
            "issue_category": ("Issue Category (ประเภทของปัญหา)", "Issue Category"),
            "issue_detail": ("Issue Detail (รายละเอียดของปัญหา)", "Issue Detail"),
            "qc_result": ("QC Result (ผลตรวจงานที่พบ)", "QC Result"),
            "root_cause_status": ("Root Cause Status (สาเหตุของปัญหาที่ยืนยันแล้ว)", "Root Cause Status"),
            "root_cause_type": ("Root Cause Type (ประเภทของสาเหตุ)", "Root Cause Type"),
            "root_cause_detail": ("Root Cause Detail (รายละเอียด)", "Root Cause Detail"),
            "immediate_action": ("Immediate Action",),
            "preventive_action": ("Preventive Action",),
            "rework": ("Rework (เข้าแก้ไข)", "Rework"),
            "service_mind": ("Service Mind (0/1)", "Service Mind"),
            "owner": ("Owner / Sub Supervisor", "Owner"),
            "case_status": ("Status (สถานะเคส)", "Status"),
            "close_date": ("Close Date (วันที่เข้าตรวจเช็ค/แก้ไข)", "Close Date"),
        },
    ),
    SheetContract(
        sheet_name="Technician Master",
        required_columns=("technician_id",),
        accepted_aliases={
            "technician_id": ("Tech ID", "Technician ID"),
            "technician_name": ("Technician Name", "Name", "Tech Name"),
            "team": ("Team",),
            "vendor": ("Vendor", "LSP", "Sub Vendor"),
            "status": ("Status",),
            "area": ("Area",),
        },
    ),
)


def find_headers(headers: Iterable[object], contract: SheetContract) -> dict[str, int]:
    indexes = {normalise_header(value): index for index, value in enumerate(headers) if value is not None}
    found: dict[str, int] = {}
    for field, aliases in contract.accepted_aliases.items():
        match = next((indexes[normalise_header(alias)] for alias in aliases if normalise_header(alias) in indexes), None)
        if match is not None:
            found[field] = match
    return found


def locate_header(sheet, contract: SheetContract, search_rows: int = 10) -> tuple[int | None, tuple[object, ...], dict[str, int]]:
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(search_rows, sheet.max_row), values_only=True), start=1
    ):
        found = find_headers(row, contract)
        if contract.required_columns[0] in found:
            return row_number, tuple(row), found
    return None, (), {}


def validate_workbook(path: str | Path) -> list[ImportValidation]:
    """Validate source structure without inferring missing operational outcomes."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    validations: list[ImportValidation] = []
    try:
        for contract in CONTRACTS:
            if contract.sheet_name not in workbook.sheetnames:
                validations.append(ImportValidation(contract.sheet_name, None, 0, contract.required_columns, 0))
                continue
            sheet = workbook[contract.sheet_name]
            header_row, _, found = locate_header(sheet, contract)
            missing = tuple(field for field in contract.required_columns if field not in found)
            rows = 0
            blank_key_rows = 0
            if header_row is not None:
                key_index = found.get(contract.required_columns[0])
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore", UserWarning)
                    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                        if not any(value not in (None, "") for value in row):
                            continue
                        rows += 1
                        if key_index is not None and (key_index >= len(row) or row[key_index] in (None, "")):
                            blank_key_rows += 1
            validations.append(ImportValidation(contract.sheet_name, header_row, rows, missing, blank_key_rows))
    finally:
        workbook.close()
    return validations
