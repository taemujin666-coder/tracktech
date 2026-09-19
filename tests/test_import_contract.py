import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from app.application.import_contract import validate_workbook
from app.application.workbook_import import preview_workbook
from app.domain.import_models import CaseLinkStatus, QCEvidenceStatus, TechnicianIdentityStatus


JOB_HEADERS = [
    "Job No.", "Install Date", "Customer", "Product / Model", "Tech ID", "Technician Name",
    "Project", "Vendor", "QC Date", "QC Result", "QC Issue Found", "Complaint", "Rework",
    "Integrity violation", "Job Status", "Note",
]
CASE_HEADERS = [
    "Project ID", "Job No. (เลขที่ออเดอร์)", "Complete Date (วันที่ติดตั้งสำเร็จ)",
    "Complaint Date (วันที่ได้รับการร้องเรียน)", "Tech ID", "Technician Name (ชื่อทีมช่าง)",
    "Issue Category (ประเภทของปัญหา)", "Issue Detail (รายละเอียดของปัญหา)",
    "QC Result (ผลตรวจงานที่พบ)", "Root Cause Status (สาเหตุของปัญหาที่ยืนยันแล้ว)",
    "Root Cause Type (ประเภทของสาเหตุ)", "Root Cause Detail (รายละเอียด)", "Immediate Action",
    "Preventive Action", "Rework (เข้าแก้ไข)", "Service Mind (0/1)", "Owner / Sub Supervisor",
    "Status (สถานะเคส)", "Close Date (วันที่เข้าตรวจเช็ค/แก้ไข)",
]
TECH_HEADERS = ["Tech ID", "Technician Name", "Team", "Vendor", "Status", "Start Date", "Product", "Area"]


def build_workbook(path: Path, case_status: str = "Resolved") -> None:
    workbook = Workbook()
    job = workbook.active
    job.title = "Job Data"
    job.append(["Job Data — title"])
    job.append(JOB_HEADERS)
    job.append(["JOB001", datetime(2026, 1, 1), "Private", "Fridge", "T001", "One", "Mall", "Vendor", datetime(2026, 1, 1), "Pass", None, "Yes", "No", "No", "Completed", None])
    job.append(["JOB001", datetime(2026, 1, 1), "Private", "Washer", "T001", "One", "Mall", "Vendor", datetime(2026, 1, 1), "Fail", None, None, None, None, "Completed", None])
    job.append(["JOB002", datetime(2026, 1, 2), "Private", "TV", "PWS1", None, "Mall", "Vendor", datetime(2026, 1, 2), "Pass", None, None, None, None, "Completed", None])
    job.append(["JOB003", datetime(2026, 1, 3), "Private", "TV", "KSP70", None, "Mall", "Vendor", datetime(2026, 1, 3), None, None, None, None, None, "Not Complete", None])
    job.append(["JOB004", datetime(2026, 1, 4), "Private", "TV", None, None, "Mall", "Vendor", None, None, None, None, None, None, "Cancel", None])

    complaint = workbook.create_sheet("Complaint Log")
    complaint.append(["Complaint Log — title"])
    complaint.append(CASE_HEADERS)
    complaint.append(["Mall", "JOB001", datetime(2026, 1, 1), datetime(2026, 1, 5), "T002", "Two", "Workmanship", "Evidence detail", "Fail", "Confirmed", "Technician", "Cause", "Fix", None, "Yes", 1, "Vendor", case_status, datetime(2026, 1, 6)])
    complaint.append(["Mall", "OLD001", datetime(2025, 12, 1), datetime(2026, 1, 7), "T001", "One", "Product", "Old evidence", "Pass", "Pending Investigation", None, None, None, None, None, 1, "Vendor", None, None])
    complaint.append(["Mall", "MISSING2026", datetime(2026, 1, 2), datetime(2026, 1, 8), "T001", "One", "Product", "Current evidence", "Pass", "Pending Investigation", None, None, None, None, None, 1, "Vendor", None, None])

    technician = workbook.create_sheet("Technician Master")
    technician.append(["Technician Master"])
    technician.append(TECH_HEADERS)
    technician.append(["T001", " One ", "A", "Vendor", "Active", None, None, "Central"])
    technician.append(["T001", "One", "A", "Vendor", "Active", None, None, "Central"])
    technician.append(["T002", "Two", "A", "Vendor", "Active", None, None, "Central"])
    workbook.save(path)


class ImportContractTest(unittest.TestCase):
    def test_real_layout_header_row_two_is_accepted(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            build_workbook(path)
            result = validate_workbook(path)
        self.assertTrue(all(item.is_valid for item in result))
        self.assertEqual({item.header_row for item in result}, {2})

    def test_missing_required_column_is_rejected(self):
        workbook = Workbook()
        workbook.active.title = "Job Data"
        workbook.active.append(["Title"])
        workbook.active.append(["Tech ID"])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "invalid.xlsx"
            workbook.save(path)
            result = validate_workbook(path)
        job_data = result[0]
        self.assertIn("job_no", job_data.missing_columns)
        self.assertFalse(job_data.is_valid)

    def test_preview_applies_evidence_first_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            build_workbook(path)
            preview = preview_workbook(path)

        self.assertEqual(preview.summary["jobs"]["rows"], 5)
        self.assertEqual(preview.summary["jobs"]["unique_job_numbers"], 4)
        self.assertEqual(preview.summary["jobs"]["duplicate_job_numbers"], 1)
        self.assertEqual(preview.summary["technicians"], {"source_rows": 3, "unique_verified": 2, "duplicate_ids": 1})
        by_job = {record.job_no: record for record in preview.jobs}
        self.assertEqual(by_job["JOB002"].technician_identity_status, TechnicianIdentityStatus.UNVERIFIED_TECHNICIAN_IDENTITY)
        self.assertIsNone(by_job["JOB002"].technician_id)
        fridge = next(record for record in preview.jobs if record.product_model == "Fridge")
        self.assertTrue(fridge.complaint)
        self.assertFalse(fridge.rework)
        self.assertFalse(fridge.integrity_violation)
        self.assertEqual(by_job["JOB003"].technician_identity_status, TechnicianIdentityStatus.PENDING_MASTER_MATCH)
        self.assertEqual(by_job["JOB003"].qc_evidence_status, QCEvidenceStatus.NOT_APPLICABLE_NOT_COMPLETE)
        self.assertEqual(by_job["JOB004"].qc_evidence_status, QCEvidenceStatus.NOT_APPLICABLE_CANCELLED)
        self.assertEqual(preview.cases[0].link_status, CaseLinkStatus.TECHNICIAN_CONFLICT)
        self.assertEqual(preview.cases[1].link_status, CaseLinkStatus.REFERENCE_OUTSIDE_CURRENT_JOB_DATA)
        self.assertEqual(preview.cases[2].link_status, CaseLinkStatus.JOB_REFERENCE_REQUIRES_REVIEW)
        self.assertEqual(preview.summary["complaints"]["reporting_date_field"], "Complaint Date")
        self.assertEqual(preview.summary["complaints"]["period_start"], "2026-01-05")
        self.assertEqual(preview.summary["complaints"]["period_end"], "2026-01-08")
        self.assertEqual(preview.summary["complaints"]["by_month"], {"2026-01": 3})
        self.assertTrue(any("รับเรื่องร้องเรียนในปี 2026" in warning for warning in preview.warnings))

    def test_case_key_is_stable_when_case_status_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "first.xlsx"
            second = Path(directory) / "second.xlsx"
            build_workbook(first, "Resolved")
            build_workbook(second, "Closed")
            before = preview_workbook(first)
            after = preview_workbook(second)
        self.assertEqual(before.cases[0].source_case_key, after.cases[0].source_case_key)
        self.assertNotEqual(before.cases[0].record_hash, after.cases[0].record_hash)


if __name__ == "__main__":
    unittest.main()
